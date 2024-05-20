import re
import os
from docx import Document

sınıflar=["mavi","beyaz","kirmizi","kahverengi","turuncu","turkuvaz","siyah","yesil","gri","pembe","sari","tanımsız"]
file_path="rapor.doc"

def count_specific_words(file_path, sınıflar):
    word_counts = {word: 0 for word in sınıflar}

    with open(file_path, 'r', encoding='utf-8') as file:
        for line in file:
            words = re.findall(r'\b\w+\b', line.lower())
            for word in words:
                if word in word_counts:
                    word_counts[word] += 1

    return word_counts

word_counts = count_specific_words(file_path, sınıflar)
#for word, count in word_counts.items():
results_list = [(word, count) for word, count in word_counts.items()]    
    #print(f"'{word}' kelimesi {count} kez geçti.")
file_path2="kaydedilen_rapor.doc"
word_count2 = count_specific_words(file_path2, sınıflar)

results_list2 = [(word, count) for word, count in word_counts.items()]
#print(results_list)
def compare_tuples(results_list, results_list2):
    comparison_results = []
    for comparison_tuple in results_list2:
        word, expected_count = comparison_tuple
        actual_count = next((count for w, count in results_list if w == word), None)
        if actual_count is not None:
            if actual_count == expected_count:
                comparison_results.append((word, actual_count, True))
            else:
                comparison_results.append((word, actual_count, False))
        else:
            comparison_results.append((word, 0, False))
    return comparison_results

# Karşılaştırma sonucunu yazdır
comparison_results = compare_tuples(results_list, results_list2)
for word, actual_count, is_match in comparison_results:
    if is_match:
        print(f"'{word}' kelimesi beklenen {actual_count} kez geçti.")
    else:
         expected_count = next((count for w, count in results_list2 if w == word), None)
         print(f"'{word}' kelimesi beklenen {expected_count} kez geçmedi, {actual_count} kez geçti.")


def raporu_doldur(file_path):
    # Belgeyi aç
    doc = Document(file_path)

    # Placeholder'ları ve değişkenleri tanımla
    placeholders = {
        "Şirket Adı": "………………..",
        "Raporlama Tarihi": "..../.…/.…",
        "Kontrol Saatleri": "..:.. / …:…",
        "Gemi Adı": "……………….",
        "Alıcı Adı": "………."
    }

    # Placeholder'ları doldur
    for paragraph in doc.paragraphs:
        for key, value in placeholders.items():
            if value in paragraph.text:
                if key == "Şirket Adı":
                    print("Şirket adını giriniz:")
                    sirket_adi = input()
                    paragraph.text = paragraph.text.replace(value, sirket_adi)
                if key == "Raporlama Tarihi":
                    print("Raporlama tarihini giriniz: ..../..../.... formatında ")
                    tarih = input()
                    paragraph.text = paragraph.text.replace(value, tarih)
                if key == "Kontrol Saatleri":
                    print("Kontrol saatlerini yazın: ..:.. / …:… formatında ")
                    saatler = input()
                    paragraph.text = paragraph.text.replace(value, saatler)
                if key == "Gemi Adı":
                    print("Gemi adı:")
                    gemi_adi = input()
                    paragraph.text = paragraph.text.replace(value, gemi_adi)
                if key == "Alıcı Adı":
                    print("Alıcı adı:")
                    alici_adi = input()
                    paragraph.text = paragraph.text.replace(value, alici_adi)
                
                





    # Düzenlenmiş belgeyi kaydet
    doc.save("Düzenlenmiş_Rapor.docx")

# Raporu doldur
raporu_doldur("Yeni Microsoft Word Belgesi.docx")
