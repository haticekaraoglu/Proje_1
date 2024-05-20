from docx import Document

def raporu_doldur(file_path):
    # Belgeyi aç
    doc = Document(file_path)

    # Placeholder'ları ve değişkenleri tanımla
    placeholders = {
        "Şirket Adı": "………………..",
        "Raporlama Tarihi": "..../.…/.…",
        "Kontrol Saatleri": "..:.. / …:…",
        "Gemi Adı": "……………….",
        "Alıcı Adı": "………."
    }

    # Placeholder'ları doldur
    for paragraph in doc.paragraphs:
        for key, value in placeholders.items():
            if value in paragraph.text:
                if key == "Şirket Adı":
                    print("Şirket adını giriniz:")
                    sirket_adi = input()
                    paragraph.text = paragraph.text.replace(value, sirket_adi)
                if key == "Raporlama Tarihi":
                    print("Raporlama tarihini giriniz: ..../..../.... formatında ")
                    tarih = input()
                    paragraph.text = paragraph.text.replace(value, tarih)
                if key == "Kontrol Saatleri":
                    print("Kontrol saatlerini yazın: ..:.. / …:… formatında ")
                    saatler = input()
                    paragraph.text = paragraph.text.replace(value, saatler)
                if key == "Gemi Adı":
                    print("Gemi adı:")
                    gemi_adi = input()
                    paragraph.text = paragraph.text.replace(value, gemi_adi)
                if key == "Alıcı Adı":
                    print("Alıcı adı:")
                    alici_adi = input()
                    paragraph.text = paragraph.text.replace(value, alici_adi)

    # Düzenlenmiş belgeyi kaydet
    doc.save("Düzenlenmiş_Rapor.docx")

# Raporu doldur
raporu_doldur("Taşımacılık Rapor Sonucu.docx")
